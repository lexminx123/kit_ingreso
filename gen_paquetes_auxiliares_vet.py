"""
gen_paquetes_auxiliares_vet.py — Genera 5 paquetes personalizados para los
auxiliares veterinarios identificados por la empresa.

Para cada auxiliar genera:
  1. Solicitud de Empleo (con datos prellenados)
  2. Contrato de Trabajo (Auxiliar Veterinario)
  3. Descripción de Cargo (Auxiliar Veterinario)
  4. Notificación de Riesgos (Auxiliar Veterinario)
  5. Carta de Aceptación General

Cada paquete se archiva en una carpeta por auxiliar y luego se empaqueta
en un ZIP único para descarga.
"""
import os, sys, shutil, zipfile
sys.path.insert(0, "/home/z/my-project/output")
from _common import *
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL

# ============================================================
# DATOS DE LOS 5 AUXILIARES VETERINARIOS
# ============================================================
AUXILIARES = [
    {
        "nombre": "Valeria Katiuska Rangel Pérez",
        "cedula": "V-27.279.527",
        "cedula_raw": "27279527",
        "direccion": "Urbanización Solar de la Quinta, Etapa 1, Los Teques, Estado Miranda",
        "telefono": "0412-983.21.38",
        "correo": "valeria.rangel02@gmail.com",
        "carpeta": "01_Valeria_Katiuska_Rangel_Perez",
    },
    {
        "nombre": "Pedro José Baldovino Cardozo",
        "cedula": "V-27.180.661",
        "cedula_raw": "27180661",
        "direccion": "San Antonio Los Altos, Rosaleda Sur, Edif. Yuruarí, Estado Miranda",
        "telefono": "0412-174.63.59",
        "correo": "baldovinop8@gmail.com",
        "carpeta": "02_Pedro_Jose_Baldovino_Cardozo",
    },
    {
        "nombre": "Cinthia Paola Hernández Padilla",
        "cedula": "V-25.562.814",
        "cedula_raw": "25562814",
        "direccion": "Las Adjuntas, Sector La Gran Parada, Casa N° 19, Estado Miranda",
        "telefono": "0414-239.76.98",
        "correo": "cinthiapaolahp@gmail.com",
        "carpeta": "03_Cinthia_Paola_Hernandez_Padilla",
    },
    {
        "nombre": "Joxelina Jacqueline Serrano",
        "cedula": "V-14.058.904",
        "cedula_raw": "14058904",
        "direccion": "Macarena Sur, Calle El Cristo, Estado Miranda",
        "telefono": "0412-928.94.96",
        "correo": "serranoyake04@gmail.com",
        "carpeta": "04_Joxelina_Jacqueline_Serrano",
    },
    {
        "nombre": "Rosa Margarita Blanco Martínez",
        "cedula": "V-14.850.506",
        "cedula_raw": "14850506",
        "direccion": "La Estrella, Sector La Mascota, Casa S/N, Calle Vargas, Estado Miranda",
        "telefono": "0412-961.30.16",
        "correo": "rosyblanco819@gmail.com",
        "carpeta": "05_Rosa_Margarita_Blanco_Martinez",
    },
]

CARGO = "AUXILIAR VETERINARIO(A)"
SALARIO_BASE = "50$"
TOTAL_MENSUAL = "250$"

# ============================================================
# 1. SOLICITUD DE EMPLEO PRELLENADA
# ============================================================
def gen_solicitud_prellenada(aux, out_path):
    doc = Document()
    section = setup_a4_portrait(doc, margins_cm=1.8)
    add_membrete(doc, "FICHA DE INGRESO", "Solicitud de Empleo", version="Versión 3.2  ·  RR.HH.")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("SOLICITUD DE EMPLEO Y FICHA DE INGRESO")
    style_run(r, size=14, bold=True, color=TEAL_DARK)
    sep = doc.add_paragraph()
    sep.paragraph_format.space_after = Pt(8)
    add_hr(sep, color="0F766E", sz="8")

    # Declaración inicial con datos prellenados
    add_para(doc,
        f"Yo, {aux['nombre']}, titular de la cédula de identidad N° {aux['cedula']}, "
        "solicito empleo en GRUPO CAVAL 1003, C.A. (ALIKA PETS) y declaro que la "
        "información aquí proporcionada es veraz, completa y exacta. Acepto que "
        "cualquier falsificación u omisión podrá ser causal de rechazo de la "
        "solicitud o de despido justificado si se descubriere con posterioridad.",
        size=10, space_after=6)

    # Tabla con datos personales prellenados
    add_section(doc, "1. DATOS PERSONALES")
    datos = [
        ("Apellidos:", aux['nombre'].split()[-2] if len(aux['nombre'].split()) > 2 else "",
         "Nombres:", " ".join(aux['nombre'].split()[:-2]) if len(aux['nombre'].split()) > 2 else aux['nombre']),
        ("Cédula de Identidad:", aux['cedula'], "Lugar de expedición:", "Venezuela"),
        ("Teléfono móvil:", aux['telefono'], "Correo electrónico:", aux['correo']),
        ("Dirección de habitación:", aux['direccion'], "", ""),
    ]
    tbl = doc.add_table(rows=len(datos), cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    for i, w in enumerate([Cm(3.5), Cm(5.0), Cm(3.5), Cm(5.0)]):
        tbl.columns[i].width = w
    for ri, (l1, v1, l2, v2) in enumerate(datos):
        bg = GRAY_ALT if ri % 2 == 0 else WHITE
        c = tbl.rows[ri].cells[0]; c.width = Cm(3.5)
        write_cell(c, l1, size=8, bold=True, color=TEAL_DARK, bg=SLATE_BG)
        c = tbl.rows[ri].cells[1]; c.width = Cm(5.0)
        write_cell(c, v1, size=9, bg=bg)
        if l2:
            c = tbl.rows[ri].cells[2]; c.width = Cm(3.5)
            write_cell(c, l2, size=8, bold=True, color=TEAL_DARK, bg=SLATE_BG)
            c = tbl.rows[ri].cells[3]; c.width = Cm(5.0)
            write_cell(c, v2, size=9, bg=bg)
        else:
            merged = tbl.rows[ri].cells[2].merge(tbl.rows[ri].cells[3])
            merged.text = ""
            set_cell_borders(merged, color="FFFFFF", sz="0")
    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(6)

    # Datos laborales (cargo ya asignado)
    add_section(doc, "2. DATOS LABORALES")
    datos_lab = [
        ("Cargo al que aplica:", "Auxiliar Veterinario(a)", "Departamento / Área:", "Clínica Veterinaria"),
        ("Salario mensual ofrecido:", f"USD {SALARIO_BASE} + bonos no salariales", "Tipo de contrato:", "Tiempo determinado (3 meses + prórroga)"),
        ("Supervisor inmediato:", "Directora Gerente / Médico(a) Veterinario(a)", "Período de prueba:", "15 días continuos"),
    ]
    tbl = doc.add_table(rows=len(datos_lab), cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    for i, w in enumerate([Cm(3.5), Cm(5.0), Cm(3.5), Cm(5.0)]):
        tbl.columns[i].width = w
    for ri, (l1, v1, l2, v2) in enumerate(datos_lab):
        bg = GRAY_ALT if ri % 2 == 0 else WHITE
        c = tbl.rows[ri].cells[0]; c.width = Cm(3.5)
        write_cell(c, l1, size=8, bold=True, color=TEAL_DARK, bg=SLATE_BG)
        c = tbl.rows[ri].cells[1]; c.width = Cm(5.0)
        write_cell(c, v1, size=9, bg=bg)
        c = tbl.rows[ri].cells[2]; c.width = Cm(3.5)
        write_cell(c, l2, size=8, bold=True, color=TEAL_DARK, bg=SLATE_BG)
        c = tbl.rows[ri].cells[3]; c.width = Cm(5.0)
        write_cell(c, v2, size=9, bg=bg)
    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(6)

    # Declaración jurada
    add_section(doc, "DECLARACIÓN JURADA")
    add_para(doc,
        "Declaro bajo fe de juramento que toda la información aquí contenida es veraz y "
        "completa. Autorizo a GRUPO CAVAL 1003, C.A. a verificar los datos suministrados, "
        "así como a consultar mis referencias personales y laborales. Acepto que la "
        "falsificación u omisión de información podrá ser causal de rechazo de la "
        "solicitud o de despido justificado.", size=9, space_after=6)

    add_para(doc,
        "En la ciudad de Los Teques, Estado Miranda, a los ____ días del mes de "
        "________________ de ________.", size=9, space_after=10)

    add_signature_block(doc, ["EL(LA) SOLICITANTE", "DIRECTORA GERENTE"])
    add_footer(section, "Solicitud de Empleo v3.2 (prellenada)")

    doc.save(out_path)
    return out_path


# ============================================================
# 2. CONTRATO DE TRABAJO PRELLENADO
# ============================================================
def gen_contrato_prellenado(aux, out_path):
    doc = Document()
    section = setup_a4_portrait(doc, margins_cm=2.0)
    add_membrete(doc, "CONTRATO INDIVIDUAL DE TRABAJO", CARGO, version="Versión 4.0  ·  RR.HH.")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"CONTRATO DE TRABAJO — {CARGO}")
    style_run(r, size=14, bold=True, color=TEAL_DARK)
    sep = doc.add_paragraph()
    sep.paragraph_format.space_after = Pt(10)
    add_hr(sep, color="0F766E", sz="8")

    # Partes
    add_para(doc,
        "Entre GRUPO CAVAL 1003, C.A., RIF N° J501662533, marca comercial ALIKA PETS, "
        "dedicada a la actividad de clínica veterinaria, tienda de mascotas y peluquería "
        "canina, con domicilio en Av. Francisco de Miranda, Local N° 1, Los Teques, "
        "Estado Miranda, en adelante «LA EMPRESA», representada en este acto por su "
        "Directora Gerente, ciudadana ESNATLIM ELENA SIMOZA, titular de la cédula de "
        f"identidad N° V-17.976.287; por una parte; y por la otra, el(la) ciudadano(a) "
        f"{aux['nombre']}, venezolano(a), mayor de edad, titular de la cédula de identidad "
        f"N° {aux['cedula']}, con domicilio en {aux['direccion']}, teléfono "
        f"{aux['telefono']} y correo electrónico {aux['correo']}, en adelante "
        "«EL(LA) TRABAJADOR(A)», han convenido en celebrar el presente Contrato de "
        "Trabajo, regido por las siguientes cláusulas:", size=10, space_after=8)

    # CLÁUSULA 1: OBJETO
    add_section(doc, "CLÁUSULA PRIMERA: OBJETO")
    add_para(doc,
        f"LA EMPRESA contrata los servicios personales de EL(LA) TRABAJADOR(A) bajo una "
        f"sola y única relación de trabajo, quien se obliga a prestar con carácter de "
        f"exclusividad sus servicios en el cargo de {CARGO}, desempeñando satisfactoriamente "
        f"las labores y responsabilidades detalladas en la Descripción de Cargo respectiva.")
    add_para(doc,
        "Las funciones principales del cargo incluyen: asistir al médico veterinario en "
        "consultas, tratamientos, laboratorios y cirugías; cuidado de animales "
        "hospitalizados; recolección de muestras biológicas; preparación y desinfección "
        "del quirófano e instrumental; gestión de citas; y mantenimiento de limpieza "
        "y bioseguridad.", space_after=6)

    # CLÁUSULA 2: DURACIÓN
    add_section(doc, "CLÁUSULA SEGUNDA: DURACIÓN DEL CONTRATO")
    add_para(doc,
        "El presente contrato se celebra por TIEMPO DETERMINADO, con duración de TRES "
        "(3) MESES, con un período de prueba de quince (15) días continuos. Podrá ser "
        "prorrogado por mutuo acuerdo. Fecha de inicio: ____ / ____ / ________. Fecha de "
        "culminación: ____ / ____ / ________.", space_after=6)

    # CLÁUSULA 3: JORNADA
    add_section(doc, "CLÁUSULA TERCERA: JORNADA DE TRABAJO")
    add_para(doc,
        "La jornada será fijada por LA EMPRESA según las necesidades del servicio. La "
        "jornada ordinaria no excederá de cinco (5) días a la semana y el horario será "
        "notificado por escrito al trabajador. LA EMPRESA se reserva el derecho de "
        "reorganizar la jornada por turnos o guardias, dentro de los límites legales.",
        space_after=6)

    # CLÁUSULA 4: SALARIO (formato nuevo sin artículos)
    add_section(doc, "CLÁUSULA CUARTA: SALARIO, BONOS Y PRESTACIONES SOCIALES")
    add_para(doc,
        f"EL(LA) TRABAJADOR(A) devengará un salario básico mensual de {SALARIO_BASE} "
        "vigente conforme a la última Gaceta Oficial relacionada con dichos conceptos. "
        "Adicionalmente, recibirá los siguientes conceptos:", size=10, space_after=4)

    # Tabla 5 conceptos
    tbl = doc.add_table(rows=6, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    widths = [Cm(6.0), Cm(4.6), Cm(6.0)]
    for i, w in enumerate(widths):
        tbl.columns[i].width = w
    headers = ["Concepto", "Monto mensual", "Naturaleza"]
    for ci, h in enumerate(headers):
        c = tbl.rows[0].cells[ci]; c.width = widths[ci]
        write_cell(c, h, size=9, bold=True, color=RGBColor(0xFF,0xFF,0xFF),
                   align=WD_ALIGN_PARAGRAPH.CENTER, bg=TEAL_HDR_BG)
    rows = [
        ("Salario básico", "50$", "Salarial (base prestaciones)"),
        ("Bono de alimentación (Cestaticket)", "80$", "No salarial (Ley Programa Alimentación Trabajador)"),
        ("Bono de transporte", "40$", "No salarial (Rembolso de gastos)"),
        ("Bono de Buen Vivir", "40$", "No salarial (Asistencial)"),
        ("Bono de Renta Telefónica", "40$", "No salarial (Rembolso de gastos)"),
    ]
    for ri, row in enumerate(rows, start=1):
        bg = GRAY_ALT if ri % 2 == 0 else WHITE
        for ci, val in enumerate(row):
            c = tbl.rows[ri].cells[ci]; c.width = widths[ci]
            write_cell(c, val, size=9, bold=(ci==0), color=BLACK, bg=bg,
                       align=WD_ALIGN_PARAGRAPH.LEFT if ci != 1 else WD_ALIGN_PARAGRAPH.CENTER)
    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(4)

    add_para(doc,
        "Las prestaciones sociales se calcularán sobre la base del salario normal "
        "devengado. LA EMPRESA depositará las prestaciones sociales mensualmente en la "
        "cuenta individual del trabajador, dentro de los primeros cinco (5) días hábiles "
        "del mes siguiente. Igualmente, LA EMPRESA cotizará al IVSS, FAOV, INCES y demás "
        "organismos de seguridad social conforme a la legislación aplicable.",
        size=10, space_after=6)

    # CLÁUSULA 5: EQUIPOS
    add_section(doc, "CLÁUSULA QUINTA: EQUIPOS, HERRAMIENTAS Y BIENES DE LA EMPRESA")
    add_para(doc,
        "LA EMPRESA pondrá a disposición de EL(LA) TRABAJADOR(A) todos los equipos médicos, "
        "instrumental clínico y quirúrgico, equipos de laboratorio, materiales, insumos, "
        "mobiliario, instalaciones y demás bienes necesarios para la correcta prestación "
        "de los servicios. Estos bienes son propiedad exclusiva de LA EMPRESA y se "
        "entregan en comodato para el ejercicio de las funciones del cargo.")
    add_para(doc,
        "EL(LA) TRABAJADOR(A) se obliga a cuidar los equipos como un buen padre de familia, "
        "reportar de inmediato cualquier daño, deterioro, falla o pérdida, no sustraer ni "
        "utilizar los equipos para fines personales o externos, mantener los equipos en "
        "condiciones de higiene y bioseguridad, y devolver todos los equipos al término de "
        "la relación laboral, en el mismo estado en que los recibió, salvo el deterioro "
        "natural por su uso adecuado. Los descuentos por pérdidas o daños se tramitarán "
        "conforme a la legislación aplicable.", space_after=6)

    # CLÁUSULA 6: OBLIGACIONES Y PROHIBICIONES
    add_section(doc, "CLÁUSULA SEXTA: OBLIGACIONES Y PROHIBICIONES")
    add_para(doc, "Son obligaciones de EL(LA) TRABAJADOR(A):", size=10, bold=True, space_after=2)
    obl = [
        "Cumplir con las órdenes e instrucciones de LA EMPRESA en todo lo concerniente al trabajo.",
        "Concurrir al trabajo en el horario establecido y permanecer en él durante la jornada.",
        "Observar las normas de bioseguridad aplicables y el uso obligatorio del EPP asignado.",
        "Mantener en buen estado los equipos, herramientas, mobiliario e instalaciones de la empresa.",
        "Guardar secreto sobre la información técnica, comercial y administrativa de LA EMPRESA.",
        "Cumplir el Reglamento Interno y el Código de Conducta de la empresa.",
        "Obtener autorización previa y por escrito para publicar en redes sociales contenido de la clínica.",
    ]
    for o in obl:
        add_bullet(doc, o, size=10)
    add_para(doc, "Son prohibiciones para EL(LA) TRABAJADOR(A):", size=10, bold=True, space_before=4, space_after=2)
    proh = [
        "Trabajar en estado de embriaguez o bajo influencia de sustancias estupefacientes o psicotrópicas.",
        "Sustraer, ocultar o dañar herramientas, mercancía, equipo, pacientes o documentos de LA EMPRESA.",
        "Abandonar el trabajo durante la jornada sin causa justificada y sin autorización del superior.",
        "Ejecutar labores por cuenta propia o de terceros dentro del horario laboral.",
        "Divulgar información confidencial, datos de clientes o historias clínicas de pacientes.",
    ]
    for pr in proh:
        add_bullet(doc, pr, size=10, color=RED_CRIT)
    add_para(doc, "", size=4, space_after=2)

    # CLÁUSULA 7: PENALIDADES
    add_section(doc, "CLÁUSULA SÉPTIMA: PENALIDADES Y RÉGIMEN DISCIPLINARIO")
    add_para(doc,
        "El incumplimiento de las obligaciones aquí previstas se sancionará conforme al "
        "siguiente régimen disciplinario simplificado de tres (3) niveles, sin perjuicio "
        "de las previsiones del Reglamento Interno:", size=10, space_after=4)

    tbl = doc.add_table(rows=4, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    widths = [Cm(1.5), Cm(7.5), Cm(7.5)]
    for i, w in enumerate(widths):
        tbl.columns[i].width = w
    headers = ["Nivel", "Tipo de falta", "Medida disciplinaria"]
    for ci, h in enumerate(headers):
        c = tbl.rows[0].cells[ci]; c.width = widths[ci]
        write_cell(c, h, size=9, bold=True, color=RGBColor(0xFF,0xFF,0xFF),
                   align=WD_ALIGN_PARAGRAPH.CENTER, bg=TEAL_HDR_BG)
    rows = [
        ("1", "LEVE — Incumplimientos menores de procedimiento, impuntualidad ocasional, omisiones leves.",
         "Amonestación VERBAL por el jefe inmediato. Se deja constancia en el expediente."),
        ("2", "GRAVE — Reincidencia de faltas leves, incumplimiento de protocolos de bioseguridad o atención al cliente, descuido de inventario.",
         "Amonestación ESCRITA firmada por el trabajador y la Directora Gerente. Se archiva en expediente."),
        ("3", "MUY GRAVE — Faltas graves: robo, abandono del trabajo, indisciplina, acoso, revelación de secretos, daño intencional a bienes de la empresa o de pacientes.",
         "DESPIDO JUSTIFICADO conforme a la legislación aplicable. Sin derecho a indemnización."),
    ]
    for ri, row in enumerate(rows, start=1):
        bg = GRAY_ALT if ri % 2 == 0 else WHITE
        for ci, val in enumerate(row):
            c = tbl.rows[ri].cells[ci]; c.width = widths[ci]
            color = RED_CRIT if ci == 2 and "DESPIDO" in val else BLACK
            write_cell(c, val, size=8, bold=(ci==0 or (ci==2 and "DESPIDO" in val)), color=color, bg=bg,
                       align=WD_ALIGN_PARAGRAPH.CENTER if ci == 0 else WD_ALIGN_PARAGRAPH.LEFT)
    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(4)

    # CLÁUSULA 8: TERMINACIÓN
    add_section(doc, "CLÁUSULA OCTAVA: TERMINACIÓN, LITIGIOS Y DOMICILIO")
    add_para(doc,
        "El presente contrato terminará por las causales previstas en la legislación "
        "aplicable. Las controversias que se susciten con motivo de la interpretación "
        "o ejecución del presente contrato serán sustanciadas ante la Inspectoría del "
        "Trabajo con competencia en el Estado Miranda, o ante la jurisdicción laboral "
        "ordinaria. Para todos los efectos legales, las partes eligen como domicilio "
        "especial y procesal la ciudad de Los Teques, Estado Miranda, a cuyos "
        "tribunales se someten expresamente.", space_after=8)

    # Cierre
    add_para(doc,
        "Se hacen dos (2) ejemplares de un mismo tenor y a un solo efecto, en la ciudad "
        "de Los Teques, a los ____ días del mes de ________________ de ________.",
        size=10, space_after=10)

    add_signature_block(doc, ["LA EMPRESA", "EL(LA) TRABAJADOR(A)"])
    add_footer(section, f"Contrato Auxiliar Veterinario v4.0 — {aux['nombre']}")

    doc.save(out_path)
    return out_path


# ============================================================
# 3. CARTA DE ACEPTACIÓN GENERAL PRELLENADA
# ============================================================
def gen_carta_aceptacion_prellenada(aux, out_path):
    doc = Document()
    section = setup_a4_portrait(doc, margins_cm=2.2)
    add_membrete(doc, "CIERRE DE EXPEDIENTE", "Aceptación general del Kit", version="Versión 3.0  ·  RR.HH.")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("CARTA DE ACEPTACIÓN GENERAL DEL KIT DE INGRESO")
    style_run(r, size=14, bold=True, color=TEAL_DARK)
    sep = doc.add_paragraph()
    sep.paragraph_format.space_after = Pt(10)
    add_hr(sep, color="0F766E", sz="8")

    add_para(doc,
        "Ciudad y fecha:  Los Teques, ____ de ________________ de ________",
        size=10, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=10)

    add_para(doc,
        "Ciudadana\nESNATLIM ELENA SIMOZA\nDirectora Gerente\nGRUPO CAVAL 1003, C.A. (ALIKA PETS)\nSu despacho.-",
        size=10, space_after=8)

    add_para(doc, "Ref: Aceptación general del Kit de Ingreso.", size=10, bold=True, space_after=8)

    add_para(doc,
        f"Yo, {aux['nombre']}, titular de la cédula de identidad N° {aux['cedula']}, "
        f"en mi condición de trabajador(a) de GRUPO CAVAL 1003, C.A. (ALIKA PETS), "
        f"ingresado(a) en el cargo de {CARGO} con fecha ____ / ____ / ________, por "
        "medio de la presente declaro:", size=10, space_after=6)

    declaraciones = [
        ("PRIMERO:  ",
         "Que he recibido copia íntegra y legible de todos los documentos que conforman el "
         "Kit de Ingreso del Trabajador, conforme al Checklist Maestro firmado por separado, "
         "incluyendo, sin carácter limitativo: contrato individual de trabajo, descripción "
         "de cargo, reglamento interno, código de conducta, política de confidencialidad, "
         "política de uso de redes sociales, autorización para depósito de prestaciones "
         "sociales, designación de beneficiarios, notificación de riesgos, hoja de recorrido "
         "habitual, acta de entrega de EPP, cartilla de bioseguridad veterinaria, protocolo "
         "de mordeduras y zoonosis, procedimiento de reporte de incidentes, autorización de "
         "tratamiento de datos personales, autorización de uso de imagen, y autorización "
         "de vigilancia por cámaras."),
        ("SEGUNDO:  ",
         "Que he leído en su totalidad cada uno de los documentos recibidos, he recibido "
         "explicaciones verbales sobre su contenido por parte de la Dirección, y he tenido "
         "la oportunidad de formular preguntas, las cuales fueron respondidas "
         "satisfactoriamente."),
        ("TERCERO:  ",
         "Que comprendo mis obligaciones, prohibiciones y deberes, y me comprometo a "
         "cumplirlos fielmente durante toda la vigencia de mi relación laboral con la "
         "empresa, así como las normas de seguridad y salud laboral establecidas."),
        ("CUARTO:  ",
         "Que entiendo que el incumplimiento de las disposiciones contenidas en los "
         "documentos del Kit podrá dar lugar a las sanciones disciplinarias previstas en "
         "el Reglamento Interno y, según la gravedad, a la terminación de la relación "
         "laboral por causa justificada."),
        ("QUINTO:  ",
         "Que autorizo expresamente a GRUPO CAVAL 1003, C.A. para el tratamiento de mis "
         "datos personales conforme a la normativa aplicable."),
        ("SEXTO:  ",
         "Que acepto que la presente firma constituye prueba fehaciente de la entrega y "
         "recepción de todos los documentos del Kit, renunciando a alegar desconocimiento "
         "de su contenido en el futuro."),
        ("SÉPTIMO:  ",
         "Que reconozco que GRUPO CAVAL 1003, C.A. ha cumplido con sus obligaciones de "
         "información, capacitación y entrega de equipos, y que estoy en condiciones de "
         "iniciar mis labores."),
    ]
    for label, txt in declaraciones:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.3
        rl = p.add_run(label)
        style_run(rl, size=10, bold=True, color=TEAL_DARK)
        r = p.add_run(txt)
        style_run(r, size=10)

    add_para(doc, "", space_after=4)
    add_para(doc,
        "En fe de lo cual firmo la presente carta en la ciudad de Los Teques, Estado "
        "Miranda, en la fecha indicada arriba.", size=10, space_after=12)

    add_signature_block(doc, ["EL(LA) TRABAJADOR(A)", "DIRECTORA GERENTE"])
    add_footer(section, f"Carta Aceptación General — {aux['nombre']}")

    doc.save(out_path)
    return out_path


# ============================================================
# GENERAR LOS 5 PAQUETES
# ============================================================
BASE_DIR = "/home/z/my-project/output/PAQUETES_AUXILIARES_VET"
os.makedirs(BASE_DIR, exist_ok=True)

# Copiar documentos base (Notificación de Riesgos y Descripción de Cargo)
NOTIF_RIESGOS_SRC = "/home/z/my-project/output/05_SEGURIDAD_LABORAL/Notificacion_Riesgos_Auxiliar_Veterinario.docx"
DESCRIPCION_CARGO_SRC = "/home/z/my-project/output/03_DESCRIPCION_DE_CARGOS/03e_Funciones_Auxiliar_Veterinario.docx"

print("Generando 5 paquetes personalizados para auxiliares veterinarios...")
print("=" * 70)

for aux in AUXILIARES:
    carpeta_aux = os.path.join(BASE_DIR, aux['carpeta'])
    os.makedirs(carpeta_aux, exist_ok=True)

    # 1. Solicitud de Empleo prellenada
    sol_path = os.path.join(carpeta_aux, "01_Solicitud_Empleo.docx")
    gen_solicitud_prellenada(aux, sol_path)

    # 2. Contrato de Trabajo prellenado
    cont_path = os.path.join(carpeta_aux, "02_Contrato_Trabajo_Auxiliar_Veterinario.docx")
    gen_contrato_prellenado(aux, cont_path)

    # 3. Descripción de Cargo (copiar el estándar)
    desc_path = os.path.join(carpeta_aux, "03_Descripcion_Cargo_Auxiliar_Veterinario.docx")
    shutil.copy2(DESCRIPCION_CARGO_SRC, desc_path)

    # 4. Notificación de Riesgos (copiar el estándar)
    notif_path = os.path.join(carpeta_aux, "04_Notificacion_Riesgos_Auxiliar_Veterinario.docx")
    shutil.copy2(NOTIF_RIESGOS_SRC, notif_path)

    # 5. Carta de Aceptación General prellenada
    carta_path = os.path.join(carpeta_aux, "05_Carta_Aceptacion_General.docx")
    gen_carta_aceptacion_prellenada(aux, carta_path)

    size_kb = sum(os.path.getsize(os.path.join(carpeta_aux, f)) for f in os.listdir(carpeta_aux)) / 1024
    print(f"  ✓ {aux['carpeta']}  ({len(os.listdir(carpeta_aux))} docs, {size_kb:.1f} KB)")

# Crear README en la carpeta base
readme_path = os.path.join(BASE_DIR, "LEEME.txt")
with open(readme_path, "w", encoding="utf-8") as f:
    f.write("PAQUETES PERSONALIZADOS — AUXILIARES VETERINARIOS\n")
    f.write("ALIKA PETS · Grupo Caval 1003, C.A.\n")
    f.write("=" * 70 + "\n\n")
    f.write("Este paquete contiene 5 carpetas, una por cada auxiliar veterinario:\n\n")
    for aux in AUXILIARES:
        f.write(f"  {aux['carpeta']}/\n")
        f.write(f"    Nombre: {aux['nombre']}\n")
        f.write(f"    Cédula: {aux['cedula']}\n")
        f.write(f"    Cargo: Auxiliar Veterinario(a)\n")
        f.write(f"    Documentos incluidos:\n")
        f.write(f"      01_Solicitud_Empleo.docx (con datos prellenados)\n")
        f.write(f"      02_Contrato_Trabajo_Auxiliar_Veterinario.docx (prellenado)\n")
        f.write(f"      03_Descripcion_Cargo_Auxiliar_Veterinario.docx\n")
        f.write(f"      04_Notificacion_Riesgos_Auxiliar_Veterinario.docx\n")
        f.write(f"      05_Carta_Aceptacion_General.docx (prellenada)\n\n")
    f.write("=" * 70 + "\n")
    f.write("Total: 25 documentos Word (.docx) + este archivo LEEME.txt\n")
    f.write("\nGenerado por: ALIKA PETS - RR.HH.\n")
    f.write("Fecha: Septiembre 2026\n")
    f.write("Versión: 4.0\n")

print("\n" + "=" * 70)
print("✓ LEEME.txt creado")

# ============================================================
# CREAR ZIP ÚNICO PARA DESCARGA
# ============================================================
ZIP_PATH = "/home/z/my-project/output/Paquetes_Auxiliares_Veterinarios.zip"
print("\nCreando ZIP único para descarga...")
with zipfile.ZipFile(ZIP_PATH, 'w', zipfile.ZIP_DEFLATED) as zipf:
    # Añadir LEEME
    zipf.write(readme_path, "LEEME.txt")
    # Añadir todos los paquetes
    for aux in AUXILIARES:
        carpeta_aux = os.path.join(BASE_DIR, aux['carpeta'])
        for f in os.listdir(carpeta_aux):
            file_path = os.path.join(carpeta_aux, f)
            arcname = os.path.join(aux['carpeta'], f)
            zipf.write(file_path, arcname)

size_mb = os.path.getsize(ZIP_PATH) / (1024 * 1024)
print(f"✓ ZIP creado: {ZIP_PATH}  ({size_mb:.2f} MB)")
print(f"\nTotal documentos en ZIP: {sum(len(os.listdir(os.path.join(BASE_DIR, aux['carpeta']))) for aux in AUXILIARES)} + LEEME.txt")
